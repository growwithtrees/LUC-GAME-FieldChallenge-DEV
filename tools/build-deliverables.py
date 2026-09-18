#!/usr/bin/env python3
"""Regenerate every printable Field Guide Challenge deliverable from this repo.

WHY THIS EXISTS
    The wall sheets, QR cards and facilitator sheet all restate data that lives in
    app.js — team numbers, slugs, box colours, fragments, reply words, and which stop
    each team starts on. Hand-maintaining them guarantees the printed pack eventually
    disagrees with the live game. Everything here is derived from app.js and
    tools/plant-wall.json, so a change to the game is one rerun away from a correct
    print pack. Never hand-edit the generated PDFs.

WHEN TO RUN
    After ANY change to the TEAMS block in app.js (numbers, colours, words, fragments,
    `start` offsets), or to tools/plant-wall.json (which plant, which photo, which organ).

USAGE
    python3 tools/build-deliverables.py              # -> the OneDrive Deliverables folder
    python3 tools/build-deliverables.py --out DIR    # -> somewhere else
    python3 tools/build-deliverables.py --only wall  # wall | qr | facilitator | hannah | roles

REQUIREMENTS
    pip install qrcode img2pdf pillow python-docx      (all installed as of 2026-09-18)
    ImageMagick (`magick`) and LibreOffice (`soffice`) on PATH.

GOTCHAS THAT COST REAL TIME — see [[gwt_print_deliverables_toolchain]]
  * img2pdf with NO layout_fun sizes each page from that image's embedded DPI, which is
    what gives landscape pages for landscape photos. Passing a fixed `pagesize` forces
    one orientation and letterboxes everything else onto a portrait page.
  * `magick -annotate` dies with "unable to read font ''" unless you pass an explicit
    -font path. There is no usable default.
  * Verify QR codes by DECODING THE RENDERED PDF, not the source images:
        pdftoppm -r 200 -png "Team QR Cards.pdf" z && zbarimg -q --raw z-*.png
    then curl each decoded URL. A wrong QR is only discoverable after printing.
  * Always render the .docx to PDF and LOOK at it. Reading the facilitator sheet as
    Hannah would is what caught reply words that were two words jammed together.
"""

import argparse, json, os, re, subprocess, sys, glob, shutil

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DEFAULT_OUT = os.path.expanduser(
    "~/Library/CloudStorage/OneDrive-GrowWithTrees/GWT Shortcuts/$GWT Info/"
    "LUC/Projects/Field Guide Challenge/Deliverables")
BASE_URL = "https://growwithtrees.github.io/LUC-GAME-FieldChallenge-DEV"
FONT_R = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_B = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
HEXC = {"Red": "#C0492F", "Green": "#2E7D46", "Blue": "#2563A8",
        "Amber": "#C6841A", "Purple": "#6D4AA8"}
STOP_NAMES = ["Messages", "Field Guide", "Scan", "Field Clips", "ChatTrain", "Field Bites"]


def read_teams():
    """Parse the TEAMS block out of app.js. Single source of truth for every artifact."""
    src = open(os.path.join(REPO, "app.js"), encoding="utf-8").read()
    block = re.search(r"var TEAMS = \{(.*?)\n  \};", src, re.S).group(1)
    pat = (r'"(\d+)":\s*\{ slug: "([\w-]+)",\s*color: "(\w+)",\s*'
           r'trainer: "(\w+)", word: "(\w+)",\s*frag: "(\d+)", start: (\d)')
    teams = [dict(zip(["name", "slug", "color", "trainer", "word", "frag", "start"], m.groups()))
             for m in re.finditer(pat, block)]
    if not teams:
        sys.exit("ERROR: could not parse TEAMS out of app.js — did its formatting change?")
    for t in teams:
        t["startname"] = STOP_NAMES[int(t["start"]) % len(STOP_NAMES)]
    return teams


def build_wall(out):
    """One page per plant, numbered, deliberately UNLABELLED, orientation-matched."""
    from PIL import Image
    import img2pdf
    plants = json.load(open(os.path.join(HERE, "plant-wall.json")))
    tmp = os.path.join(HERE, "_wall"); os.makedirs(tmp, exist_ok=True)
    missing = [p["src"] for p in plants if not os.path.exists(os.path.expanduser(p["src"]))]
    if missing:
        sys.exit("ERROR: wall source image(s) not found:\n  " + "\n  ".join(missing))
    for p in plants:
        src = os.path.expanduser(p["src"])
        w, h = Image.open(src).size
        land = w > h
        canvas = "1650x1275" if land else "1275x1650"      # letter at 150dpi
        fit = "1560x1185" if land else "1185x1560"          # 0.3in margins, maximised
        subprocess.run(["magick", src, "-auto-orient", "-resize", fit,
                        "-background", "white", "-gravity", "center", "-extent", canvas,
                        "-font", FONT_R, "-pointsize", "80", "-fill", "#444444",
                        "-gravity", "southeast", "-annotate", "+55+45", str(p["n"]),
                        "-density", "150", "-units", "PixelsPerInch", "-quality", "94",
                        os.path.join(tmp, "%02d.jpg" % p["n"])], check=True)
        p["orient"] = "landscape" if land else "portrait"
    files = sorted(glob.glob(os.path.join(tmp, "*.jpg")))
    dest = os.path.join(out, "2 Print - Plant Wall")
    os.makedirs(os.path.join(dest, "Individual Sheets"), exist_ok=True)
    # NO layout_fun: each page takes its own image's orientation. See module docstring.
    with open(os.path.join(dest, "Plant Wall - Print Sheets.pdf"), "wb") as f:
        f.write(img2pdf.convert(files))
    for old in glob.glob(os.path.join(dest, "Individual Sheets", "*.jpg")):
        os.remove(old)
    tag = {p["n"]: p.get("tag", "") for p in plants}
    for p in plants:
        extra = " (%s)" % tag[p["n"]] if tag[p["n"]] else ""
        shutil.copy(os.path.join(tmp, "%02d.jpg" % p["n"]),
                    os.path.join(dest, "Individual Sheets",
                                 "%02d %s - %s%s [%s].jpg"
                                 % (p["n"], p["name"].split(" /")[0], p["organ"], extra, p["orient"])))
    json.dump(plants, open(os.path.join(HERE, "_wallkey.json"), "w"), indent=1)
    shutil.rmtree(tmp)
    print("  wall: %d sheets" % len(files))


def build_qr(teams, out):
    """One card per team: number, box colour, QR, and the plain URL as a fallback."""
    import qrcode, img2pdf
    from PIL import Image, ImageDraw, ImageFont
    tmp = os.path.join(HERE, "_qr"); os.makedirs(tmp, exist_ok=True)
    W, H = 1275, 825                                   # half-letter at 150dpi
    for i, t in enumerate(teams, 1):
        url = "%s/%s/" % (BASE_URL, t["slug"])
        q = qrcode.QRCode(box_size=10, border=2,
                          error_correction=qrcode.constants.ERROR_CORRECT_M)
        q.add_data(url); q.make(fit=True)
        qi = q.make_image(fill_color="black", back_color="white").convert("RGB") \
              .resize((560, 560), Image.NEAREST)
        card = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(card)
        d.rectangle([0, 0, W, 18], fill=HEXC[t["color"]])
        card.paste(qi, (70, 150)); x = 700
        d.text((x, 160), "Team", font=ImageFont.truetype(FONT_R, 52), fill="#6b7280")
        d.text((x, 215), t["name"], font=ImageFont.truetype(FONT_B, 150), fill="#111111")
        d.ellipse([x, 400, x + 38, 438], fill=HEXC[t["color"]])
        d.text((x + 56, 397), "%s box" % t["color"],
               font=ImageFont.truetype(FONT_B, 44), fill=HEXC[t["color"]])
        d.text((x, 485), "Scan to start your quest.",
               font=ImageFont.truetype(FONT_R, 38), fill="#374151")
        d.text((x, 537), "One phone can drive; everyone",
               font=ImageFont.truetype(FONT_R, 32), fill="#6b7280")
        d.text((x, 577), "helps solve.", font=ImageFont.truetype(FONT_R, 32), fill="#6b7280")
        d.text((70, 745), url, font=ImageFont.truetype(FONT_R, 26), fill="#9ca3af")
        card.save(os.path.join(tmp, "%02d.jpg" % i), quality=92, dpi=(150, 150))
    pages = []
    for i in range(0, len(teams), 2):
        p = Image.new("RGB", (1275, 1650), "white")
        p.paste(Image.open(os.path.join(tmp, "%02d.jpg" % (i + 1))), (0, 0))
        if i + 2 <= len(teams):
            p.paste(Image.open(os.path.join(tmp, "%02d.jpg" % (i + 2))), (0, 825))
        ImageDraw.Draw(p).line([(0, 824), (1275, 824)], fill="#cccccc", width=2)
        fn = os.path.join(tmp, "page%d.jpg" % (i // 2 + 1))
        p.save(fn, quality=92, dpi=(150, 150)); pages.append(fn)
    lay = img2pdf.get_layout_fun(pagesize=(img2pdf.in_to_pt(8.5), img2pdf.in_to_pt(11)),
                                 fit=img2pdf.FitMode.into)
    dest = os.path.join(out, "3 Print - Team QR Cards"); os.makedirs(dest, exist_ok=True)
    with open(os.path.join(dest, "Team QR Cards.pdf"), "wb") as f:
        f.write(img2pdf.convert(pages, layout_fun=lay))
    shutil.rmtree(tmp)
    print("  qr: %d cards" % len(teams))
    print("  VERIFY BEFORE PRINTING: pdftoppm -r 200 -png 'Team QR Cards.pdf' z && "
          "zbarimg -q --raw z-*.png   (then curl each URL)")


def build_roles(out):
    """Scanner / Navigator / Scribe, three to a page. Print 10x and cut."""
    import img2pdf
    from PIL import Image, ImageDraw, ImageFont
    ROLES = [
        ("SCANNER", "#2E7D46", "Holds the phone.",
         "You drive the Let Grow app — Identify, search, video, chat.\n"
         "Before you scan, set what you are photographing:\n"
         "leaf, flower, fruit or bark. Wrong setting, bad match."),
        ("NAVIGATOR", "#2563A8", "Reads the clue out loud.",
         "You keep the team pointed at the right task. Read the whole\n"
         "clue before anyone touches the app — the detail that\nmatters is usually in the last line."),
        ("SCRIBE", "#C6841A", "Runs the team's page.",
         "You type the codes in and keep the page open. Do not close\n"
         "it — it remembers where you are. One phone drives the whole\nquest, so keep it charged."),
    ]
    W, H = 1275, 550
    page = Image.new("RGB", (1275, 1650), "white")
    for i, (name, hexc, tag, body) in enumerate(ROLES):
        c = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(c)
        d.rectangle([0, 0, 22, H], fill=hexc)
        d.text((70, 60), name, font=ImageFont.truetype(FONT_B, 64), fill=hexc)
        d.text((70, 140), tag, font=ImageFont.truetype(FONT_B, 30), fill="#111111")
        y = 200
        for line in body.split("\n"):
            d.text((70, y), line, font=ImageFont.truetype(FONT_R, 27), fill="#374151"); y += 40
        d.text((70, H - 70), "Field Guide Challenge",
               font=ImageFont.truetype(FONT_R, 22), fill="#9ca3af")
        page.paste(c, (0, i * 550))
    d = ImageDraw.Draw(page)
    for y in (550, 1100):
        d.line([(0, y), (1275, y)], fill="#cccccc", width=2)
    tmp = os.path.join(HERE, "_roles.jpg")
    page.save(tmp, quality=92, dpi=(150, 150))
    dest = os.path.join(out, "5 Print - Hannah and Roles"); os.makedirs(dest, exist_ok=True)
    with open(os.path.join(dest, "Role Cards.pdf"), "wb") as f:
        f.write(img2pdf.convert([tmp]))
    os.remove(tmp)
    print("  roles: 3 cards, 1 page (print 10x and cut)")


def _docx_to_pdf(path, dest):
    subprocess.run(["/Applications/LibreOffice.app/Contents/MacOS/soffice", "--headless",
                    "--convert-to", "pdf", "--outdir", dest, path],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def build_facilitator(teams, out):
    """Lindsey's 4-page sheet. Delegates to _facilitator_sheet.py, which holds the layout."""
    dest = os.path.join(out, "4 Print - Facilitator"); os.makedirs(dest, exist_ok=True)
    json.dump(teams, open(os.path.join(HERE, "_teams.json"), "w"), indent=1)
    subprocess.run([sys.executable, os.path.join(HERE, "_facilitator_sheet.py")],
                   cwd=HERE, check=True)
    src = os.path.join(HERE, "Facilitator Sheet.docx")
    shutil.copy(src, os.path.join(dest, "Facilitator Sheet.docx"))
    _docx_to_pdf(src, dest)
    os.remove(src)
    print("  facilitator: Facilitator Sheet.docx + .pdf")


def build_hannah(teams, out):
    """Hannah's own one-page card. She must NOT be given the facilitator sheet —
    it carries every answer in the game and her only job is replying one word."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HEAD, BODY = "Plus Jakarta Sans", "Inter"
    INK, MUTE, GREEN = RGBColor(0x11, 0x11, 0x11), RGBColor(0x6b, 0x72, 0x80), RGBColor(0x2E, 0x7D, 0x46)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.5)
    st = doc.styles["Normal"]; st.font.name = BODY; st.font.size = Pt(11); st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.15
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

    def shade(cell, hexc):
        el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), hexc)
        cell._tc.get_or_add_tcPr().append(el)

    def para(t, size=11, color=INK, bold=False, italic=False, after=5, before=0):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(t); r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold; r.font.italic = italic
        return p

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Hannah — Messages desk"); r.font.name = HEAD; r.font.size = Pt(24)
    r.font.bold = True; r.font.color.rgb = GREEN
    r._element.rPr.rFonts.set(qn('w:eastAsia'), HEAD)
    para("Field Guide Challenge · Tree Line USA · this is your whole job", 10, MUTE,
         italic=True, after=12)
    para("Ten teams will message you during the game. Each one sends:", 11, after=6)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    r = p.add_run("TEAM 4 CHECKING IN"); r.font.name = "Courier New"; r.font.size = Pt(15)
    r.font.bold = True
    para("Read the number. Find it below. Reply with that team's ONE WORD and nothing else. "
         "That word is what unlocks their next stop.", 11, after=10)

    half = (len(teams) + 1) // 2
    t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
    for i, htxt in enumerate(["They send", "You reply", "They send", "You reply"]):
        c = t.rows[0].cells[i]; shade(c, "1C5230"); c.width = Inches(1.85)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        rr = c.paragraphs[0].add_run(htxt); rr.font.bold = True; rr.font.size = Pt(10)
        rr.font.name = HEAD; rr.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    for i in range(half):
        cells = t.add_row().cells
        pair = [teams[i]] + ([teams[i + half]] if i + half < len(teams) else [])
        for k, tm in enumerate(pair):
            for j, val in enumerate(["TEAM %s" % tm["name"], tm["word"].upper()]):
                c = cells[k * 2 + j]; c.width = Inches(1.85)
                c.paragraphs[0].paragraph_format.space_after = Pt(3)
                rr = c.paragraphs[0].add_run(val); rr.font.size = Pt(13); rr.font.bold = (j == 1)
                if j == 1:
                    rr.font.name = "Courier New"

    para("", 6, after=6)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("If something goes odd"); r.font.name = HEAD; r.font.size = Pt(14)
    r.font.bold = True; r.font.color.rgb = INK
    for k, v in [
        ("A team sends something else", "Reply: “Say again — send me your team name and CHECKING IN.”"),
        ("Two or three arrive at once", "Just answer each in turn. They are not racing each other."),
        ("You get swamped", "Tell Lindsey. The fallback is one shared word for everybody — she can announce it and the game keeps running."),
        ("A team says the word did not work", "Check you sent the right row. Words are single words, no spaces."),
    ]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(k + " — "); r1.font.bold = True; r1.font.size = Pt(10.5)
        r2 = p.add_run(v); r2.font.size = Pt(10.5)
    para("Teams start at different stops on purpose, so the messages will trickle in across the "
         "whole game rather than all at the start. Stay on your phone until Lindsey calls time.",
         10, MUTE, before=10)

    dest = os.path.join(out, "5 Print - Hannah and Roles"); os.makedirs(dest, exist_ok=True)
    src = os.path.join(HERE, "Hannah Card.docx")
    doc.save(src)
    shutil.copy(src, os.path.join(dest, "Hannah Card.docx"))
    _docx_to_pdf(src, dest)
    os.remove(src)
    print("  hannah: Hannah Card.docx + .pdf")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=DEFAULT_OUT, help="deliverables folder")
    ap.add_argument("--only", choices=["wall", "qr", "facilitator", "hannah", "roles"],
                    help="build just one artifact")
    a = ap.parse_args()
    teams = read_teams()
    print("Read %d teams from app.js: %s" % (len(teams), ", ".join(t["name"] for t in teams)))
    os.makedirs(a.out, exist_ok=True)
    want = (lambda k: a.only in (None, k))
    if want("wall"): build_wall(a.out)
    if want("qr"): build_qr(teams, a.out)
    if want("facilitator"): build_facilitator(teams, a.out)
    if want("hannah"): build_hannah(teams, a.out)
    if want("roles"): build_roles(a.out)
    print("Done -> %s" % a.out)


if __name__ == "__main__":
    main()
