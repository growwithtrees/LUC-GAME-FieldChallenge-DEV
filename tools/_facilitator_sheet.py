import json, docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAMS = json.load(open("_teams.json"))
HEAD, BODY = "Plus Jakarta Sans", "Inter"
INK, MUTE, GREEN = RGBColor(0x11,0x11,0x11), RGBColor(0x6b,0x72,0x80), RGBColor(0x2E,0x7D,0x46)
HEX = {"Red":"C0492F","Green":"2E7D46","Blue":"2563A8","Amber":"C6841A","Purple":"6D4AA8"}

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.5)
st = doc.styles["Normal"]; st.font.name = BODY; st.font.size = Pt(10); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.12
st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

def shade(cell, hexcolor):
    el = OxmlElement('w:shd'); el.set(qn('w:val'),'clear'); el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

def h(text, size=20, color=INK, before=10, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.font.name = HEAD; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn('w:eastAsia'), HEAD)
    return p

def para(text, size=10, color=INK, italic=False, bold=False, after=4, space_before=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(space_before)
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color
    r.font.italic = italic; r.font.bold = bold
    return p

def table(headers, rows, widths, header_fill="1C5230"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (c, w) in enumerate(zip(t.rows[0].cells, widths)):
        c.width = Inches(w); shade(c, header_fill)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        r = c.paragraphs[0].add_run(headers[i]); r.font.bold = True; r.font.size = Pt(9)
        r.font.name = HEAD; r.font.color.rgb = RGBColor(0xff,0xff,0xff)
    for row in rows:
        cells = t.add_row().cells
        for i, (c, val) in enumerate(zip(cells, row)):
            c.width = Inches(widths[i])
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            txt, bold, fill = (val if isinstance(val, tuple) else (val, False, None))
            if fill: shade(c, fill)
            r = c.paragraphs[0].add_run(txt); r.font.size = Pt(9); r.font.bold = bold
            if fill: r.font.color.rgb = RGBColor(0xff,0xff,0xff)
    return t

# ── header ────────────────────────────────────────────────────────────────
h("Field Guide Challenge — Facilitator Sheet", 22, GREEN, before=0, after=2)
para("Liberty Utilities Central · Tree Line USA, October 2026 · keep this sheet face down",
     9, MUTE, italic=True, after=10)

# ── the teams ─────────────────────────────────────────────────────────────
h("1. Teams, boxes and combos", 14)
para("Ten numbered teams in five colour pairs — 1+2, 3+4, 5+6, 7+8, 9+10. Hand the QR cards out IN "
     "PAIRS; if you end up with fewer than ten teams, drop a whole colour so every box still has two.", 9.5, after=4)
para("Every team does all six stops, but they START on different ones and wrap around, so the room "
     "is not all messaging Hannah at once and not all at the plant wall at once. A team on their "
     "\u201cfirst\u201d stop may well be at ChatTrain. The last column is where each team begins.",
     9.5, after=6)
rows = []
for t in TEAMS:
    rows.append([("Team "+t["name"], True, None), ("●  " + t["color"], False, HEX[t["color"]]),
                 t["frag"], t["word"].upper(), (t["startname"], True, None)])
table(["Team", "Box", "Their number", "Hannah replies", "They START at"],
      rows, [1.05, 1.05, 1.15, 1.55, 2.7])
para("", 4, after=2)
combos = " · ".join(f"{TEAMS[i]['color']} {TEAMS[i]['frag']}{TEAMS[i+1]['frag']}" for i in range(0, 10, 2))
para("Box combinations:  " + combos, 10, bold=True, after=2)
para("The first team in each pair puts their number first. Set the physical locks to these before "
     "the room opens, and test each one.", 9, MUTE, after=8)

# ── the six stops ─────────────────────────────────────────────────────────
h("2. The six stops — answers, and how to unstick a team", 14)
para("Numbered in canonical order, which is NOT the order most teams meet them. Teams type the "
     "answer into their own page; if one has been stuck a few minutes, give the nudge, not the "
     "answer. The answer column is your backstop only.", 9.5, after=6)

STOPS = [
 ("1", "Messages",
  "Message Hannah “TEAM <n> CHECKING IN”; she replies one word.",
  "Their word (column 4 above)",
  "“Have you actually sent the message yet?” If Hannah is swamped, just read them their word."),
 ("2", "Field Guide · Search",
  "Low compatible shrub, coral-pink berry clusters, deer browse it hard. Name it.",
  "CORALBERRY  (buckbrush also works)",
  "“Search the field guide for the berry colour, not the plant.” Crews who learned it as buckbrush are right — the app just files it under coralberry."),
 ("3", "Plant Identify · Scan",
  "On the wall, find the shrub with its feet in the water that carries a creamy white pincushion flowerhead. Identify it (setting the flower/fruit/leaf/bark selector correctly), open its entry, read its MAX HEIGHT in feet.",
  "12",
  "First ask what they set the selector to — most bad matches are a leaf setting on a flower photo. Then: “Look for the flower that is a perfect white ball, not a flat cluster.” That is sheet 1. If they scanned elderberry (sheet 9) they will still see 12, so send them to the pincushion. Red osier dogwood means they scanned sheet 3; it reads 10 ft, so the game bounces them."),
 ("4", "Field Clips",
  "Play “What’s the Tree.” Johnny reads one tree part by part and never names it. Name it.",
  "CALLERY PEAR  (bradford pear, or just pear)",
  "“Watch it to the end — he never says it. What did the fruit look like?” Then: “It’s on the plant wall, and it’s one of the ones you’d take out.”"),
 ("5", "ChatTrain",
  "Talk Mrs. Webb down about spraying. Pass the beat to get the phrase.",
  "COMPATIBLE COVER",
  "“Tell her what you are NOT killing.” She wants to hear the work is selective and her wildlife brush stays. The AI will re-ask twice; they cannot get permanently stuck."),
 ("6", "Field Bites",
  "Finish “What Makes a Plant Compatible?” and enter the keyword on the last card.",
  "BIOLOGICAL CONTROL",
  "“You have to answer all three checks — the keyword won't show until you do.” That is deliberate; they cannot swipe to the end."),
]
table(["#", "Feature", "What they have to do", "Answer", "If they’re stuck"],
      [[(s[0], True, None), s[1], s[2], (s[3], True, None), s[4]] for s in STOPS],
      [0.3, 1.15, 2.3, 1.45, 2.3])
para("Answers are not case sensitive and ignore punctuation.", 9, MUTE, space_before=4, after=8)

doc.add_page_break()

# ── plant wall ────────────────────────────────────────────────────────────
h("3. Plant wall key", 14, before=0)
para("Ten sheets, one plant each, numbered but deliberately NOT labelled — naming them would give "
     "away stops 3 and 4. This is the only key.", 9.5, after=4)
para("The photos are deliberately a MIX of flower, fruit, leaf and bark, because Identify asks the "
     "learner which one they are photographing. Getting that selector right is the skill the stop is "
     "really teaching — a team that leaves it on the wrong setting will get poor matches and think "
     "the app is broken. The \u201cPhoto is a\u201d column is what they should be selecting. Pages match "
     "each photo's orientation so every print is as big as the paper allows.", 9.5, after=6)
WALL = [(str(w["n"]), w["name"], w["organ"], w["comp"], w["role"]) for w in json.load(open("_wallkey.json"))]
table(["Sheet", "What it is", "Photo is a", "Comp?", "Why it's on the wall"],
      [[(w[0], True, None), (w[1], True, None), (w[2], True, None), w[3], w[4]] for w in WALL],
      [0.5, 1.75, 0.9, 0.8, 3.55])
para("Teams are numbered, so nothing on this wall can be mistaken for a team's own name.", 9, MUTE, space_before=4, after=8)

# ── Hannah ────────────────────────────────────────────────────────────────
h("4. Hannah's job (stop 1)", 14)
para("Every team messages Hannah “TEAM 1 CHECKING IN”, “TEAM 2 CHECKING IN” and so on. She reads the "
     "number and replies with that team's one word, nothing else.", 9.5, after=4)
para("   ".join(f"TEAM {t['name']} → {t['word'].upper()}" for t in TEAMS[:5]), 9.5, after=2)
para("   ".join(f"TEAM {t['name']} → {t['word'].upper()}" for t in TEAMS[5:]), 9.5, after=4)
para("Wrong message? Reply: “Say again — send me your team name and CHECKING IN.” If ten teams at "
     "once turns out to be too much, just reply the same word to everybody and tell the room the "
     "word — the game still works, it only stops teams overhearing each other.", 9, MUTE, after=8)

# ── if it goes wrong ──────────────────────────────────────────────────────
h("5. If something goes wrong", 14)
for problem, fix in [
  ("A team's QR won't scan", "Type the address on the bottom of their card straight into the browser. It is short on purpose."),
  ("They closed the page / phone died", "Reopening the same link puts them back where they were — progress is saved on that phone. If they switch phones they start over, so keep one phone driving."),
  ("The app won't identify a printed plant", "Nine times out of ten the flower/fruit/leaf/bark selector is on the wrong setting — check that first. If it still will not match, tell them the species; stop 3 only needs the height off the entry."),
  ("ChatTrain is slow", "Normal — it thinks for a few seconds a turn. Let them keep typing; it will land."),
  ("A team is miles behind", "Give them the answer to a stop and move them on. Finishing together matters more than every team solving every clue."),
  ("Only one team of a pair finishes", "Give them their partner's number so they can open the box. Nobody should be stood next to a locked box at the end."),
]:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(problem + " — "); r.font.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(fix); r2.font.size = Pt(9.5)

para("Everything the teams need is inside the Let Grow app. If the app is down, the game is down — "
     "worth a check the morning of.", 9, MUTE, space_before=8)

doc.save("Facilitator Sheet.docx")
print("built Facilitator Sheet.docx")
